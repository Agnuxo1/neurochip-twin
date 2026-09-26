"""EPA NFA well trajectories, DNT references, and published hit-call replication.

Run from the workspace root with ``PYTHONPATH=data/.pylibs;repo/src``. Upstream
EPA endpoint fit tables are fetched from a pinned GitHub commit and SHA256
checked. The baseline recomputes chemical hit sums from those endpoint calls;
it does not refit tcplfit2 concentration-response curves.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
import pandas as pd

EPA_COMMIT = "01adf3e1a0068c87fe221d60df36b9f96c4b4b1d"
EPA_BASE = f"https://raw.githubusercontent.com/USEPA/CompTox-DNT-NFA-Refinement/{EPA_COMMIT}/"
UPSTREAM = {
    "DIV12_analysis_tcplfit2_res_02May2025.xlsx": (
        "source_files/", "1fe6b420033cef660b713dad02c7bbddb383f0ca18ad3119749be54b3cf07aff"),
    "AUC_analysis_tcplfit2_res_30Apr2025.xlsx": (
        "source_files/", "045c298212e4e5e833a69b79ebbad8b5433014a76a5007b25c48703b3d2dbab9"),
    "dnt_ref_tbl2.xlsx": (
        "source_files/", "0ef673dd117f77a15915c914245244b93556121e69ef09fdec9216170487652e"),
    "spid_chem_hitsum_tbl_all_3_datasets_use_this.csv": (
        "source_files/", "c3e5a34ce4293e3e862443795ff1e3590427b871048aebf115540d8d2b20c19d"),
    "Bioactivity_bin_tbl_comp_methods.csv": (
        "output/", "c0ced486478c2d33a0cccb32a3dc1ec81f2752f37bb1b7c7a515185c661fea68"),
    "RF_importance_AUC_29feb26.csv": (
        "output/", "0fa2d92ffb666052c195cfd6fb40f8baf43ca7b1c44de3a2f4721cf8389a10d7"),
    "Compare_ref_chems_23mar26_edit.xlsx": (
        "source_files/", "7aa976761e7cda2bbf5363941caf0971ab475a96c1608e004dced4586861cf1a"),
    "Results_sens_spec_BA_refine_ref2.csv": (
        "output/", "4209b59473b2d4b438335f9e528c3fdcad86fe6b9c85c18a8d267958749c93e6"),
}

# Rdata column -> EPA endpoint name. These are the 17 raw MEA readouts.
FEATURES = {
    "meanfiringrate": "firing_rate_mean",
    "burst.per.min": "burst_rate",
    "mean.isis": "per_burst_interspike_interval",
    "per.spikes.in.burst": "per_burst_spike_percent",
    "mean.dur": "burst_duration_mean",
    "mean.IBIs": "interburst_interval_mean",
    "nAE": "active_electrodes_number",
    "nABE": "bursting_electrodes_number",
    "ns.n": "network_spike_number",
    "ns.peak.m": "network_spike_peak",
    "ns.durn.m": "spike_duration_mean",
    "ns.percent.of.spikes.in.ns": "per_network_spike_spike_percent",
    "ns.mean.insis": "inter_network_spike_interval_mean",
    "ns.durn.sd": "network_spike_duration_std",
    "ns.mean.spikes.in.ns": "per_network_spike_spike_number_mean",
    "r": "correlation_coefficient_mean",
    "mi": "mutual_information_norm",
}
LABELS = {"Positive", "Negative"}
SEED = 20260926


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ensure_upstream(cache_dir: Path) -> None:
    cache_dir.mkdir(parents=True, exist_ok=True)
    for name, (folder, expected) in UPSTREAM.items():
        target = cache_dir / name
        if target.exists() and _sha256(target) == expected:
            continue
        temp = target.with_name(target.name + ".part")
        try:
            with urllib.request.urlopen(EPA_BASE + folder + name, timeout=90) as response, temp.open("wb") as output:
                for chunk in iter(lambda: response.read(1024 * 1024), b""):
                    output.write(chunk)
            if _sha256(temp) != expected:
                raise ValueError(f"EPA upstream SHA256 mismatch: {name}")
            temp.replace(target)
        finally:
            temp.unlink(missing_ok=True)


def _clean_id(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    return text if re.fullmatch(r"DTXSID\d+", text) else None


def _clean_label(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip().capitalize()
    return text if text in LABELS else None


def _name_key(value: str) -> str:
    return " ".join(str(value).casefold().split())


def _unique_map(rows: pd.DataFrame, key: str, value: str) -> dict[str, str]:
    pairs = defaultdict(set)
    for a, b in rows[[key, value]].itertuples(index=False, name=None):
        if pd.notna(a) and pd.notna(b):
            pairs[_name_key(a)].add(str(b))
    return {k: next(iter(v)) for k, v in pairs.items() if len(v) == 1}


def chemical_ids(raw_dir: Path, cache_dir: Path) -> tuple[dict[str, str], dict[str, str]]:
    local = pd.read_excel(raw_dir / "dnt_ref_tbl.xlsx")
    upstream = pd.read_csv(cache_dir / "spid_chem_hitsum_tbl_all_3_datasets_use_this.csv")
    remote = upstream[["spid", "dsstox_substance_id.div12"]].drop_duplicates()
    check = local[["spid", "dsstox_substance_id"]].merge(remote, on="spid", validate="one_to_one")
    if len(check) != 255 or not (check["dsstox_substance_id"] == check["dsstox_substance_id.div12"]).all():
        raise ValueError("Local and EPA SPID-to-DTXSID maps differ")
    primary = dict(zip(local.spid, local.dsstox_substance_id))
    ref = pd.read_excel(cache_dir / "dnt_ref_tbl2.xlsx")
    annotation = pd.read_csv(raw_dir / "annotate_dnt_ref_chems.csv")
    names = pd.concat([
        ref[["chnm", "dsstox_substance_id"]],
        annotation[["chnm", "dsstox_substance_id"]],
    ], ignore_index=True)
    secondary = _unique_map(names, "chnm", "dsstox_substance_id")
    return primary, secondary


def _read_rdata(path: Path) -> dict:
    """Read an .Rdata file with `rdata` (MIT licence). pyreadr (AGPL-3.0) is only an optional fallback."""
    try:
        import rdata
        objects = rdata.read_rda(str(path))
        return {k: (v if isinstance(v, pd.DataFrame) else pd.DataFrame(v)) for k, v in objects.items()}
    except ImportError:
        try:
            import pyreadr
        except ImportError as error:
            raise RuntimeError("Install rdata (pip install rdata) to read the EPA .Rdata file") from error
        return pyreadr.read_r(str(path))


def load_wells(raw_dir: Path, cache_dir: Path) -> pd.DataFrame:
    objects = _read_rdata(raw_dir / "All_DIV_Data.Rdata")
    source = objects["rval.dat"].copy()
    for col in source.columns:                       # R factors -> plain values
        if isinstance(source[col].dtype, pd.CategoricalDtype):
            source[col] = source[col].astype(object)
    required = set(FEATURES) | {"treatment", "apid.short", "date", "well", "DIV", "dose", "units", "spid", "srcf", "wllq_notes"}
    if not required <= set(source) or set(source.DIV) != {5, 7, 9, 12} or set(source.units) != {"uM"}:
        raise ValueError("Unexpected EPA All_DIV_Data schema or units")
    primary, secondary = chemical_ids(raw_dir, cache_dir)
    wells = source.rename(columns={
        "treatment": "chemical", "apid.short": "plate", "dose": "concentration_uM",
        "DIV": "div", "srcf": "source_file", "wllq_notes": "imputation_note", **FEATURES,
    }).copy()
    wells["dtxsid"] = wells.spid.map(primary)
    missing = wells.dtxsid.isna()
    wells.loc[missing, "dtxsid"] = wells.loc[missing, "chemical"].map(
        lambda value: secondary.get(_name_key(value)))
    wells["dtxsid"] = wells.dtxsid.map(_clean_id)
    wells["id_source"] = np.where(wells.spid.isin(primary), "spid_map", "exact_name_annotation")
    wells.loc[wells.dtxsid.isna(), "id_source"] = "unmapped"
    wells["well_type"] = np.where(wells.concentration_uM.eq(0), "vehicle", "treated")
    wells["was_imputed"] = wells.imputation_note.ne("")
    key = ["plate", "date", "div", "well", "spid", "concentration_uM", "source_file"]
    if wells.duplicated(key).any():
        raise ValueError("Duplicate physical well, chemical, and source key")
    wells = wells.sort_values(key, kind="stable").reset_index(drop=True)
    return wells


def normalize_vehicle(wells: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    groups = ["plate", "date", "div"]
    features = list(FEATURES.values())
    vehicles = wells[wells.well_type.eq("vehicle")].copy()
    physical_well = groups + ["well"]
    duplicates = vehicles[vehicles.duplicated(physical_well, keep=False)]
    if not duplicates.empty:
        inconsistent = duplicates.groupby(physical_well)[features].nunique(dropna=False).gt(1).any(axis=1)
        if inconsistent.any():
            raise ValueError("Duplicate physical vehicle wells have conflicting MEA features")
    vehicles = vehicles.drop_duplicates(physical_well)
    counts = vehicles.groupby(groups, dropna=False).size().rename("n_vehicle_wells")
    med = vehicles.groupby(groups, dropna=False)[features].median().add_prefix("vehicle_median_")
    stats = counts.to_frame().join(med)
    for feature in features:
        center = f"vehicle_median_{feature}"
        centered = vehicles.set_index(groups)[feature] - stats[center]
        mad = centered.abs().groupby(level=groups).median() * 1.4826
        stats[f"vehicle_mad_{feature}"] = mad
    stats = stats.reset_index()
    if (stats.n_vehicle_wells < 2).any():
        raise ValueError("Plate/date/DIV with fewer than two vehicle wells")
    result = wells.merge(stats, on=groups, how="left", validate="many_to_one")
    if result.n_vehicle_wells.isna().any():
        raise ValueError("A plate/date/DIV has no vehicle controls")
    for feature in features:
        delta = result[feature] - result[f"vehicle_median_{feature}"]
        scale = result[f"vehicle_mad_{feature}"].replace(0, np.nan)
        result[f"z_vehicle_{feature}"] = delta / scale
    result = result.drop(columns=[c for c in result if c.startswith("vehicle_median_") or c.startswith("vehicle_mad_")])
    return result, stats


def _evidence_map(table: pd.DataFrame, id_col: str, label_col: str) -> dict[str, str]:
    evidence = defaultdict(set)
    for identifier, label in table[[id_col, label_col]].itertuples(index=False, name=None):
        identifier, label = _clean_id(identifier), _clean_label(label)
        if identifier and label:
            evidence[identifier].add(label)
    conflict = {identifier: labels for identifier, labels in evidence.items() if len(labels) > 1}
    if conflict:
        raise ValueError(f"Conflicting labels within source: {list(conflict)[:3]}")
    return {identifier: next(iter(labels)) for identifier, labels in evidence.items()}


def _efsa_ids(path: Path) -> set[str]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book["Full dataset"]
    rows = sheet.iter_rows(min_col=4, max_col=4, values_only=True)
    next(rows)
    ids = {_clean_id(row[0]) for row in rows}
    book.close()
    return ids - {None}


def load_labels(wells: pd.DataFrame, raw_dir: Path, cache_dir: Path) -> pd.DataFrame:
    primary = pd.read_excel(cache_dir / "dnt_ref_tbl2.xlsx")
    shafer = pd.read_excel(raw_dir / "dnt_ref_tbl.xlsx")
    oecd = pd.read_excel(raw_dir / "dnt_ref_oecd_culbreth_curate.xlsx")
    sept = pd.read_excel(raw_dir / "DNT_ref_list_17sept2025.xlsx")
    efsa = _efsa_ids(raw_dir / "EFSA_Mundy_2024_list_DNT%20IVB.xlsx")
    maps = {
        "epa_ref_tbl2": _evidence_map(primary, "dsstox_substance_id", "ref"),
        "shafer_ref_tbl": _evidence_map(shafer, "dsstox_substance_id", "DNT Reference"),
        "oecd_appendix_a": _evidence_map(oecd, "DTXSID", "Appendix A"),
        "sept2025_reference": _evidence_map(sept, "DTXSID", "Reference"),
    }
    chemicals = wells[["spid", "chemical", "dtxsid", "id_source"]].drop_duplicates()
    if chemicals.spid.duplicated().any():
        raise ValueError("One SPID has multiple chemical identities")
    for source, lookup in maps.items():
        chemicals[source] = chemicals.dtxsid.map(lookup)
    chemicals["efsa_listed"] = chemicals.dtxsid.isin(efsa)
    chemicals["label"] = chemicals.epa_ref_tbl2
    chemicals["label_provenance"] = np.where(chemicals.label.notna(),
                                              "USEPA/source_files/dnt_ref_tbl2.xlsx", None)
    return chemicals.sort_values("spid").reset_index(drop=True)


def _hit_sums(cache_dir: Path) -> tuple[pd.DataFrame, dict[str, list[str]]]:
    div12 = pd.read_excel(cache_dir / "DIV12_analysis_tcplfit2_res_02May2025.xlsx")
    auc = pd.read_excel(cache_dir / "AUC_analysis_tcplfit2_res_30Apr2025.xlsx")
    importance = pd.read_csv(cache_dir / "RF_importance_AUC_29feb26.csv")
    for table, n_assays in ((div12, 7), (auc, 17)):
        if set(table.hitc.dropna()) != {0, 1} or table.groupby("name").assay.nunique().ne(n_assays).any():
            raise ValueError("Unexpected EPA endpoint hit calls or assay count")
        if table.duplicated(["name", "assay"]).any():
            raise ValueError("Duplicate endpoint hit call")
    a = div12.groupby("name").hitc.sum().rename("DIV12")
    b = auc.groupby("name").hitc.sum().rename("AUC")
    sums = pd.concat([a, b], axis=1, join="inner")
    top_features = {}
    for k in (1, 2, 3, 5):
        features = importance.Feature.head(k).str.replace("AUC", "CCTE_Shafer_MEA_dev", regex=False).tolist()
        selected = auc[auc.assay.isin(features)]
        if selected.assay.nunique() != k:
            raise ValueError(f"Top-{k} endpoint not found in AUC fit table")
        sums[f"Top{k}"] = selected.groupby("name").hitc.sum()
        top_features[f"Top{k}"] = features
    if len(sums) != 255 or sums.isna().any().any():
        raise ValueError("Expected 255 EPA fitted chemicals across DIV12 and AUC")
    return sums.reset_index(names="spid"), top_features


def reproduce_baseline(labels: pd.DataFrame, raw_dir: Path, cache_dir: Path) -> dict:
    sums, top_features = _hit_sums(cache_dir)
    fitted = sums.merge(labels[["spid", "dtxsid", "label"]], on="spid", how="left", validate="one_to_one")
    fitted = fitted[fitted.label.isin(LABELS)]
    counts = fitted.label.value_counts()
    if counts.get("Positive") != 86 or counts.get("Negative") != 19:
        raise ValueError(f"Unexpected EPA published reference intersection: {counts.to_dict()}")
    published = pd.read_csv(raw_dir / "Results_sens_spec_BA_refine.csv").set_index("Model")
    precomputed = pd.read_csv(cache_dir / "Bioactivity_bin_tbl_comp_methods.csv").set_index("spid")
    comparisons = {
        "DIV12": "div12.hitsum", "AUC": "auc.hitsum", "Top1": "top1.feature.hitsum",
        "Top2": "top2.feature.hitsum", "Top3": "top3.feature.hitsum", "Top5": "top5.feature.hitsum",
    }
    joined = sums.set_index("spid").join(precomputed[list(comparisons.values())])
    mismatches = {column: int((joined[model] != joined[column]).sum()) for model, column in comparisons.items()}
    if any(mismatches.values()):
        raise ValueError(f"EPA published hit sums differ: {mismatches}")
    models = {"DIV12": ("DIV12", 1), "AUC.3hit": ("AUC", 3),
              "AUC.1hit": ("AUC", 1), **{f"Top{k}": (f"Top{k}", 1) for k in (1, 2, 3, 5)}}
    def score(frame: pd.DataFrame, target: pd.DataFrame) -> dict:
        scored = {}
        for name, (column, threshold) in models.items():
            positive = frame.label.eq("Positive")
            prediction = frame[column].ge(threshold)
            tp = int((positive & prediction).sum())
            tn = int((~positive & ~prediction).sum())
            n_pos, n_neg = int(positive.sum()), int((~positive).sum())
            sensitivity = round(100 * tp / n_pos, 1)
            specificity = round(100 * tn / n_neg, 1)
            balanced = round(50 * (tp / n_pos + tn / n_neg), 1)
            expected = target.loc[name]
            delta = {key: round(value - float(expected[field]), 1) for key, value, field in (
                ("sensitivity_pp", sensitivity, "Sensitivity"),
                ("specificity_pp", specificity, "Specificity"),
                ("balanced_accuracy_pp", balanced, "Balanced_accuracy"),
            )}
            scored[name] = dict(threshold_hits=threshold, tp=tp, fn=n_pos - tp,
                                tn=tn, fp=n_neg - tn, n_positive=n_pos, n_negative=n_neg,
                                sensitivity_pct=sensitivity, specificity_pct=specificity,
                                balanced_accuracy_pct=balanced, published_delta_pp=delta)
        return scored

    results = score(fitted, published)
    ref2_table = pd.read_excel(cache_dir / "Compare_ref_chems_23mar26_edit.xlsx")
    ref2_labels = _evidence_map(ref2_table, "dsstox_substance_id", "ref")
    fitted_ref2 = sums.merge(labels[["spid", "dtxsid"]], on="spid", how="left", validate="one_to_one")
    fitted_ref2["label"] = fitted_ref2.dtxsid.map(ref2_labels)
    fitted_ref2 = fitted_ref2[fitted_ref2.label.isin(LABELS)]
    ref2_counts = fitted_ref2.label.value_counts()
    if ref2_counts.get("Positive") != 102 or ref2_counts.get("Negative") != 18:
        raise ValueError(f"Unexpected revised EPA reference intersection: {ref2_counts.to_dict()}")
    published_ref2 = pd.read_csv(cache_dir / "Results_sens_spec_BA_refine_ref2.csv").set_index("Model")
    ref2_results = score(fitted_ref2, published_ref2)
    return dict(method="Sum EPA tcplfit2 endpoint hitc by SPID, then threshold chemical hit count",
                endpoint_fit_stage="precomputed by EPA; concentration-response fitting not rerun",
                reference_revision="dnt_ref_tbl2",
                reference_source="USEPA/source_files/dnt_ref_tbl2.xlsx; DTXSID intersection with 255 fitted SPIDs",
                reference_revision_note="The pinned manuscript script now reads Compare_ref_chems_23mar26_edit.xlsx and writes Results_sens_spec_BA_refine_ref2.csv. The supplied Results_sens_spec_BA_refine.csv and Bioactivity_bin_tbl_comp_methods.csv match the earlier dnt_ref_tbl2 intersection used here.",
                published_comparison="data/neurotox_probe/nfa_refine/Results_sens_spec_BA_refine.csv",
                upstream_commit=EPA_COMMIT, n_fitted=255, n_reference=105,
                n_positive=86, n_negative=19, hit_sum_mismatches=mismatches,
                top_features=top_features, models=results,
                reference_sensitivity=dict(revision="Compare_ref_chems_23mar26_edit",
                                           published_comparison="USEPA/output/Results_sens_spec_BA_refine_ref2.csv",
                                           n_reference=len(fitted_ref2), n_positive=102, n_negative=18,
                                           models=ref2_results),
                comparison_explanation="All selected published differences are 0.0 percentage points: the same EPA endpoint calls, RF ranking, and dnt_ref_tbl2 reference intersection are used. Other reference lists or refitted concentration-response curves can change these metrics.",
                interpretation="Published Table 3 is in-sample. EPA RF feature ranking was not reselected within held-out folds.")


def compare_vehicle_stats(stats: pd.DataFrame, wells: pd.DataFrame, raw_dir: Path) -> dict:
    """Quantify differences from EPA's separately filtered DIV12 z-score table."""
    columns = ["date", "apid.short", "DIV", "acnm", "median.veh.controls"]
    epa = pd.read_csv(raw_dir / "rval_zval_table.csv", usecols=columns)
    n_epa_per_feature = int(epa.acnm.value_counts().iloc[0])
    epa = epa.drop_duplicates().rename(columns={"apid.short": "plate", "DIV": "div"})
    joined = epa.merge(stats, on=["plate", "date", "div"], how="left", validate="many_to_one")
    if joined.n_vehicle_wells.isna().any():
        raise ValueError("EPA z table has plate/date/DIV absent from raw well data")
    own = joined.apply(lambda row: row[f"vehicle_median_{row['acnm']}"], axis=1)
    equal = np.isclose(own.to_numpy(), joined["median.veh.controls"].to_numpy())
    return dict(raw_div12_rows=int(wells["div"].eq(12).sum()),
                epa_z_rows_per_feature=n_epa_per_feature,
                raw_div12_plate_date_groups=int(stats["div"].eq(12).sum()),
                epa_z_plate_date_groups=int(len(epa) / len(FEATURES)),
                endpoint_plate_medians_compared=len(equal),
                endpoint_plate_medians_equal=int(equal.sum()),
                note="Our vehicle medians use all raw wells, with duplicate physical vehicles counted once; EPA's z table has fewer DIV12 rows and groups.")


def frozen_splits(labels: pd.DataFrame) -> dict:
    """Assign one fold per connected component of DTXSID and normalized name."""
    rows = labels.reset_index(drop=True)
    parent = list(range(len(rows)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def union(i, j):
        parent[find(i)] = find(j)

    seen = {}
    for i, row in rows.iterrows():
        for token in (f"id:{row.dtxsid}" if pd.notna(row.dtxsid) else None,
                      f"name:{_name_key(row.chemical)}"):
            if token is None:
                continue
            if token in seen:
                union(i, seen[token])
            else:
                seen[token] = i
    members = defaultdict(list)
    for i in range(len(rows)):
        members[find(i)].append(i)
    components = []
    for indices in members.values():
        subset = rows.iloc[indices]
        known = set(subset.label.dropna())
        if len(known) > 1:
            raise ValueError("Conflicting DNT labels within a chemical group")
        components.append((sorted(indices), next(iter(known)) if known else "Unknown"))
    components.sort(key=lambda item: rows.iloc[item[0][0]].spid)
    strata = [label for _, label in components]
    if min(Counter(strata).values()) < 5:
        raise ValueError("Insufficient groups to stratify 5 folds")
    assignment = {}
    rng = np.random.default_rng(SEED)
    for stratum in sorted(set(strata)):
        indices = np.asarray([i for i, value in enumerate(strata) if value == stratum])
        rng.shuffle(indices)
        for position, component_idx in enumerate(indices):
            fold = position % 5
            for row_idx in components[component_idx][0]:
                assignment[rows.iloc[row_idx].spid] = fold
    entries = [dict(spid=row.spid, chemical=row.chemical,
                    dtxsid=None if pd.isna(row.dtxsid) else row.dtxsid,
                    label=None if pd.isna(row.label) else row.label,
                    fold=assignment[row.spid]) for row in rows.itertuples(index=False)]
    if len(entries) != rows.spid.nunique() or len({entry["spid"] for entry in entries}) != len(entries):
        raise ValueError("SPID split assignment is not one-to-one")
    fold_counts = []
    for fold in range(5):
        subset = [entry for entry in entries if entry["fold"] == fold]
        fold_counts.append(dict(fold=fold, n_spids=len(subset),
                                n_positive=sum(x["label"] == "Positive" for x in subset),
                                n_negative=sum(x["label"] == "Negative" for x in subset)))
    return dict(seed=SEED, n_folds=5, n_chemical_groups=len(components),
                group_rule="Connected by DTXSID or normalized chemical name",
                fold_semantics="fold is the held-out test fold; train on all other folds",
                fold_counts=fold_counts, assignments=entries)


def process(raw_dir: Path, out_dir: Path, results_dir: Path) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    results_dir.mkdir(parents=True, exist_ok=True)
    cache = out_dir / "upstream"
    ensure_upstream(cache)
    wells = load_wells(raw_dir, cache)
    wells, stats = normalize_vehicle(wells)
    labels = load_labels(wells, raw_dir, cache)
    baseline = reproduce_baseline(labels, raw_dir, cache)
    baseline["normalization_comparison"] = compare_vehicle_stats(stats, wells, raw_dir)
    baseline["local_input_sha256"] = {
        name: _sha256(raw_dir / name) for name in (
            "All_DIV_Data.Rdata", "dnt_ref_tbl.xlsx", "dnt_ref_oecd_culbreth_curate.xlsx",
            "DNT_ref_list_17sept2025.xlsx", "EFSA_Mundy_2024_list_DNT%20IVB.xlsx",
            "Results_sens_spec_BA_refine.csv",
        )
    }
    splits = frozen_splits(labels)
    wells.to_parquet(out_dir / "wells.parquet", index=False)
    stats.to_parquet(out_dir / "vehicle_stats.parquet", index=False)
    labels.to_parquet(out_dir / "dnt_labels.parquet", index=False)
    (results_dir / "epa_baseline_repro.json").write_text(json.dumps(baseline, indent=2) + "\n", encoding="utf-8")
    (results_dir / "splits_epa_nfa.json").write_text(json.dumps(splits, indent=2) + "\n", encoding="utf-8")
    summary = dict(n_wells=len(wells), n_spids=labels.spid.nunique(),
                   n_chemicals=labels.dtxsid.nunique(), n_plate_date_div=len(stats),
                   n_vehicle_rows=int(wells.well_type.eq("vehicle").sum()),
                   n_vehicle_unique_wells=int(stats.n_vehicle_wells.sum()),
                   n_missing_dtxsid=int(wells.dtxsid.isna().sum()),
                   baseline={name: row["balanced_accuracy_pct"] for name, row in baseline["models"].items()})
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=Path("data/neurotox_probe/nfa_refine"))
    parser.add_argument("--out-dir", type=Path, default=Path("data/processed/epa_nfa"))
    parser.add_argument("--results-dir", type=Path, default=Path("repo/results"))
    args = parser.parse_args()
    print(json.dumps(process(args.raw_dir, args.out_dir, args.results_dir), indent=2))


if __name__ == "__main__":
    main()
